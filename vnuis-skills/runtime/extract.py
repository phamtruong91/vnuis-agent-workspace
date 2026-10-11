# -*- coding: utf-8 -*-
"""Trích chữ từ tệp người dùng để đưa vào skill: .docx .xlsx .csv .tsv .txt .md .pdf (pdf cần lệnh pdftotext). Bảng Word/Excel được giữ dạng bảng dòng '|'.
Tệp quét ảnh (pdf scan) không có chữ thì cần OCR riêng, gói này không làm.
CLI: python3 extract.py tệp1 [tệp2 ...]  -> in văn bản gộp, kèm ước tính số token."""
import subprocess, sys
from pathlib import Path

def _docx(p):
    from docx import Document
    d = Document(p); out = []
    body = d.element.body
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for ch in body.iterchildren():
        if ch.tag.endswith("}p"):
            t = Paragraph(ch, d).text.strip()
            if t: out.append(t)
        elif ch.tag.endswith("}tbl"):
            for row in Table(ch, d).rows:
                out.append("| " + " | ".join(c.text.strip().replace("\n", " ") for c in row.cells) + " |")
    return "\n".join(out)

def _xlsx(p):
    from openpyxl import load_workbook
    wb = load_workbook(p, data_only=True, read_only=True); out = []
    for ws in wb.worksheets:
        out.append(f"## Sheet: {ws.title}")
        for r in ws.iter_rows(values_only=True):
            if any(c not in (None, "") for c in r): out.append("| " + " | ".join("" if c is None else str(c) for c in r) + " |")
    return "\n".join(out)

def _pdf(p):
    r = subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True, text=True)
    if r.returncode != 0: raise RuntimeError("không đọc được pdf (cần pdftotext): " + r.stderr[:200])
    if len(r.stdout.strip()) < 50: raise RuntimeError(f"{p}: gần như không có chữ, có thể là bản quét ảnh, cần OCR")
    return r.stdout

def read_file(path):
    p = Path(path); ext = p.suffix.lower()
    if ext == ".docx": return _docx(p)
    if ext in (".xlsx", ".xlsm"): return _xlsx(p)
    if ext == ".pdf": return _pdf(p)
    if ext in (".txt", ".md", ".csv", ".tsv", ".json"): return p.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"chưa hỗ trợ định dạng {ext}")

def merge(paths):
    """Gộp nhiều tệp, mỗi tệp có dòng tiêu đề để model biết nguồn."""
    return "\n\n".join(f"=== TỆP: {Path(p).name} ===\n{read_file(p)}" for p in paths)

def est_tokens(text, chars_per_token=3.0):
    """Ước lượng thô (tiếng Việt khoảng 2,5–3,5 ký tự/token tùy bộ tách từ); không phải số đo thật."""
    return int(len(text) / chars_per_token)

def chunk(text, max_chars=24000, overlap=600):
    """Chia văn bản dài theo ranh giới đoạn/dòng."""
    if len(text) <= max_chars: return [text]
    parts, cur = [], ""
    for line in text.split("\n"):
        if len(cur) + len(line) + 1 > max_chars and cur:
            parts.append(cur); cur = cur[-overlap:] if overlap else ""
        cur += line + "\n"
    if cur.strip(): parts.append(cur)
    return parts

if __name__ == "__main__":
    t = merge(sys.argv[1:]); print(t); print(f"\n[ước tính ~{est_tokens(t)} token, {len(t)} ký tự]", file=sys.stderr)
