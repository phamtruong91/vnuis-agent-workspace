# -*- coding: utf-8 -*-
"""Xuất Markdown do model soạn ra Word (.docx) hoặc Excel (.xlsx). Cần: pip install python-docx openpyxl (chỉ khi dùng xuất file)."""
import re

def _runs(par, text):
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            r = par.add_run(part); r.bold = (i % 2 == 1)

def _table_rows(lines):
    rows = []
    for l in lines:
        cells = [c.strip() for c in l.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c): continue
        rows.append(cells)
    return rows

def split_blocks(md):
    """Trả danh sách ('h', cấp, chữ) | ('p', chữ) | ('li', chữ) | ('table', hàng)."""
    out, buf = [], []
    def flush():
        nonlocal buf
        if buf: out.append(("table", _table_rows(buf))); buf = []
    for line in md.splitlines():
        if line.strip().startswith("|"): buf.append(line); continue
        flush()
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m: out.append(("h", len(m.group(1)), m.group(2).strip())); continue
        m = re.match(r"^\s*(?:[-*]|\d+[.)])\s+(.*)", line)
        if m: out.append(("li", m.group(1))); continue
        if line.strip(): out.append(("p", line.strip()))
    flush()
    return out

def to_docx(md, path):
    from docx import Document
    d = Document()
    for b in split_blocks(md):
        if b[0] == "h": d.add_heading(b[2], level=min(b[1], 4))
        elif b[0] == "p": _runs(d.add_paragraph(), b[1])
        elif b[0] == "li": _runs(d.add_paragraph(style="List Bullet"), b[1])
        else:
            rows = b[1]
            if not rows: continue
            n = max(len(r) for r in rows)
            t = d.add_table(rows=len(rows), cols=n); t.style = "Table Grid"
            for i, r in enumerate(rows):
                for j in range(n):
                    t.cell(i, j).text = r[j] if j < len(r) else ""
                    if i == 0:
                        for p in t.cell(i, j).paragraphs:
                            for run in p.runs: run.bold = True
    d.save(path); return path

def to_xlsx(md, path):
    """Mỗi bảng Markdown thành một sheet; phần chữ còn lại vào sheet 'Ghi chú'."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(); wb.remove(wb.active)
    notes, title, n = [], "Bảng", 0
    for b in split_blocks(md):
        if b[0] == "h": title = re.sub(r"[\[\]:*?/\\]", "", b[2])[:28] or "Bảng"; notes.append(b[2])
        elif b[0] in ("p", "li"): notes.append(b[1])
        elif b[1]:
            n += 1; ws = wb.create_sheet(f"{n}. {title}"[:31])
            for r in b[1]: ws.append(r)
            for c in ws[1]: c.font = Font(bold=True)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(60, max(10, max(len(str(c.value or "")) for c in col) + 2))
    ws = wb.create_sheet("Ghi chú")
    for t in notes: ws.append([t])
    ws.column_dimensions["A"].width = 120
    wb.save(path); return path

def export(md, fmt, path_base):
    if fmt == "docx": return to_docx(md, path_base + ".docx")
    if fmt == "xlsx": return to_xlsx(md, path_base + ".xlsx")
    open(path_base + ".md", "w", encoding="utf-8").write(md); return path_base + ".md"
